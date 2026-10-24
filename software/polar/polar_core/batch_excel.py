"""Portable Excel export and shared workbook specification for batch results.

The analysis remains Python-only on the user's computer. Install openpyxl using
requirements.txt for local Excel export. All formulas refer to recorded results;
changing volume/calibration requires rerunning the image analysis.
"""
from pathlib import Path


def workbook_spec(report: dict) -> dict:
    """Clean results only; detailed measurements remain in batch_report.json."""
    rows=[]
    for selected in report['selections']:
        result=selected.get('result')
        rows.append([selected['well'],result['contact_angle_mean_deg'] if result else None,
                     selected['uncertainty']['combined_direction_height_deg'] if result else None])
    return {'sheets':[{'name':'Results','rows':rows,
        'header':['Well','Contact angle (°)','Error (°)'],'header_row':1,'cells':[],
        'widths':[16,25,20],'decimals':{'B':'0.000','C':'0.000'},'last_row':len(rows)+1}]}


def export_excel(report: dict, path: Path):
    """Used by the downloadable Python batch command on any computer."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError as error:
        raise ImportError('Excel export needs openpyxl. Run: python3 -m pip install -r requirements.txt') from error
    spec=workbook_spec(report)
    workbook=Workbook()
    workbook.remove(workbook.active)
    for entry in spec['sheets']:
        sheet=workbook.create_sheet(entry['name'])
        sheet.sheet_view.showGridLines=False
        def put(cell,value):
            if isinstance(value,dict):
                cell.value=value['formula']
            else:
                cell.value=value
                # Filenames and source strings must never become Excel formulas.
                if isinstance(value,str):
                    cell.data_type='s'
            cell.font=Font(name='Arial',size=10,color='243247')
            cell.alignment=Alignment(vertical='center')
        for address,value in entry['cells']:
            put(sheet[address],value)
        for i,row in enumerate([entry['header']]+entry['rows'],entry['header_row']):
            for j,value in enumerate(row,1):
                cell=sheet.cell(i,j)
                put(cell,value)
                if i==entry['header_row']:
                    cell.fill=PatternFill('solid',fgColor='263C58')
                    cell.font=Font(name='Arial',size=10,bold=True,color='FFFFFF')
                    cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
                elif i%2==0:
                    cell.fill=PatternFill('solid',fgColor='F0F4F8')
            sheet.row_dimensions[i].height=42 if i==entry['header_row'] else 23
        for j,width in enumerate(entry['widths'],1):
            sheet.column_dimensions[get_column_letter(j)].width=width
        for column,fmt in entry['decimals'].items():
            for i in range(entry['header_row']+1,entry['header_row']+len(entry['rows'])+1):
                sheet[f'{column}{i}'].number_format=fmt
        sheet.freeze_panes=f'B{entry["header_row"]+1}'
        end=get_column_letter(len(entry['header']))
        sheet.auto_filter.ref=f'A{entry["header_row"]}:{end}{entry["header_row"]+len(entry["rows"])}'
        if entry['name']=='Summary':
            sheet['A2'].font=Font(name='Arial',size=16,bold=True,color='243247')
            sheet.row_dimensions[2].height=28
            sheet['F4'].number_format='0.000000'
            for i in range(entry['notes_start'],entry['last_row']+1):
                sheet.row_dimensions[i].height=22
            sheet.sheet_properties.tabColor='263C58'
    path.parent.mkdir(parents=True,exist_ok=True)
    workbook.save(path)
