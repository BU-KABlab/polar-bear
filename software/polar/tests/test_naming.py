import io
import json
import math
import time
import zipfile

import cv2
import numpy as np
import pytest
from openpyxl import load_workbook

from app import clean_settings,create_app
from polar_core.naming import filename_metadata,preview_names,select_timepoints
from tests.test_contact_angle import synthetic


def settings(**kw):
    return clean_settings(dict(filename_format='custom',filename_pattern='*plate.{well}_t={time}s_z={z}mm*',grouping='timepoints',**kw))


def test_configurable_pattern_preserves_decimal_height_and_time():
    m=filename_metadata('camera_123_plate.c4_t=120s_z=59-300mm_b25_20260903-183148.tiff',settings())
    assert m['well']=='C4' and m['zg_mm']==59.3 and m['elapsed_s']==120
    assert m['brightness_group']=='b25'


def test_custom_order_and_minute_conversion():
    s=clean_settings(dict(filename_format='custom',filename_pattern='{well}_height{z}_after{time}min.tif',time_unit='min',grouping='timepoints'))
    m=filename_metadata('A12_height59.3_after0.5min.tif',s)
    assert m['elapsed_s']==30 and m['zg_mm']==59.3


def test_scan_index_interval_and_origin():
    s=settings(time_kind='index',time_interval=2,time_unit='min',first_scan=1)
    assert filename_metadata('plate.C4_t=1s_z=59-300mm.tif',s)['elapsed_s']==0
    assert filename_metadata('plate.C4_t=3s_z=59-300mm.tif',s)['elapsed_s']==240
    with pytest.raises(ValueError,match='precedes'):filename_metadata('plate.C4_t=0s_z=59-300mm.tif',s)
    with pytest.raises(ValueError,match='whole'):filename_metadata('plate.C4_t=1.5s_z=59-300mm.tif',s)


@pytest.mark.parametrize('pattern',['*{unknown}*','*{well}_{well}_{z}mm*','*{well}{z}mm*','*{well}_z={z}*','*plate.{well}_t={time}s*'])
def test_invalid_patterns_rejected(pattern):
    with pytest.raises(ValueError):clean_settings(dict(filename_format='custom',filename_pattern=pattern,grouping='timepoints'))


def test_default_does_not_guess_time_from_capture_timestamp():
    s=clean_settings({})
    m=filename_metadata('camera_305_ursa_plate.C4_z=59-300mm_b25_20260903-183148.tiff',s)
    assert 'elapsed_s' not in m
    with pytest.raises(ValueError,match='containing'):clean_settings(dict(grouping='timepoints'))


def test_preview_reports_unmatched_names_and_groups():
    r=preview_names(['plate.C1_t=0s_z=59-300mm.tif','plate.C1_t=0s_z=59-500mm.tif','plate.C1_t=60s_z=59-300mm.tif','unmatched.tif'],settings())
    assert r['matched']==3 and r['group_count']==2
    assert r['rows'][-1]['status']=='unmatched'


def test_selection_never_crosses_well_or_time_and_retains_failed_groups():
    rows=[]
    for well,t,target in [('C1',0,59.5),('C1',10,59.3),('C2',0,59.7)]:
        for z in [59.3,59.5,59.7]:
            rows.append(dict(well=well,elapsed_s=t,zg_mm=z,focus_score=3 if z==target else 1,status='valid',filename=f'{well}-{t}-{z}',brightness_group='b25',ellipse_axis_ratio=1))
    rows.append(dict(well='C2',elapsed_s=10,zg_mm=59.3,status='rejected'))
    chosen=select_timepoints(rows,.2)
    assert [(r['well'],r['elapsed_s']) for r in chosen]==[('C1',0),('C1',10),('C2',0),('C2',10)]
    assert [r['zg_mm'] for r in chosen[:3]]==[59.5,59.3,59.7]
    assert chosen[-1]['status']=='failed'


def test_timepoint_zip_through_real_app_and_export(tmp_path,monkeypatch):
    monkeypatch.delenv('POLAR_API_TOKEN',raising=False)
    app=create_app(tmp_path,testing=True);client=app.test_client();headers={'X-POLAR-Request':'1'}
    try:
        image,_=synthetic();archive=io.BytesIO()
        with zipfile.ZipFile(archive,'w') as z:
            for well,t in [('C1',0),('C1',60),('C2',0),('C2',60)]:
                for h,blur in [(59.3,1.2),(59.5,.1),(59.7,2.)]:
                    pixels=cv2.GaussianBlur(image,(0,0),blur)
                    _,encoded=cv2.imencode('.png',(np.clip(pixels,0,1)*65535).astype(np.uint16))
                    z.writestr(f'images/plate.{well}_t={t}s_z={h:.3f}mm.png',encoded.tobytes())
            # A time point with no valid LED ring must remain present with a blank result.
            _,blank=cv2.imencode('.png',np.zeros((360,700),dtype=np.uint8))
            z.writestr('images/plate.C2_t=120s_z=59.500mm.png',blank.tobytes())
        archive.seek(0)
        response=client.post('/api/jobs',data={'files':(archive,'series.zip'),'settings':json.dumps(settings())},headers=headers)
        assert response.status_code==202,response.json
        jid=response.json['id']
        for _ in range(300):
            job=client.get('/api/jobs/'+jid).json
            if job['state'] not in ('queued','running'):break
            time.sleep(.02)
        assert job['state']=='complete',job
        assert len(job['results'])==5
        for result in job['results'][:4]:
            assert result['zg']==59.5 and result['angle'] is not None
        assert job['results'][-1]['angle'] is None and job['results'][-1]['elapsed_s']==120
        wb=load_workbook(io.BytesIO(client.get(f'/api/jobs/{jid}/results.xlsx').data))
        assert [c.value for c in wb.active[1]]==['Well','Time (s)','Contact angle (°)','Error (°)']
        assert wb.active.max_column==4 and wb.active.max_row==6
        assert wb.active['B2'].value==0 and wb.active['B3'].value==60
        assert wb.active['C6'].value is None
        assert job['settings']['filename_pattern']==settings()['filename_pattern']
    finally:app.extensions['polar']['pool'].shutdown(wait=True)


def test_upload_rejects_mismatched_custom_names_and_mixed_times_in_well_mode(tmp_path,monkeypatch):
    monkeypatch.delenv('POLAR_API_TOKEN',raising=False)
    app=create_app(tmp_path,testing=True);client=app.test_client();headers={'X-POLAR-Request':'1'}
    try:
        response=client.post('/api/jobs',data={'files':(io.BytesIO(b'fake'),'wrong.tif'),'settings':json.dumps(settings())},headers=headers)
        assert response.status_code==400 and 'do not match' in response.json['error']
        s=settings();s['grouping']='wells'
        response=client.post('/api/jobs',data={'files':[(io.BytesIO(b'fake'),'plate.C1_t=0s_z=59.3mm.tif'),(io.BytesIO(b'fake'),'plate.C1_t=60s_z=59.3mm.tif')],'settings':json.dumps(s)},headers=headers)
        assert response.status_code==400 and 'Multiple time points' in response.json['error']
        response=client.post('/api/naming/preview',json={'settings':settings(),'filenames':['plate.C1_t=0s_z=59.3mm.tif']},headers=headers)
        assert response.status_code==200 and response.json['matched']==1
    finally:app.extensions['polar']['pool'].shutdown(wait=True)

def test_elapsed_unit_suffix_mismatch_rejected():
    with pytest.raises(ValueError,match='time-unit suffix'):settings(time_unit='min')


def test_wellplate_height_read_directly_with_reference_z_calibration():
    name='camera_43_ursa_wellplate.B1_PAMA_water_after_dispense_white_height=124-800mm_b25_20261006-205557.tiff'
    # Legacy offset input must not silently shift filename heights in new runs.
    s=clean_settings(dict(filename_format='custom',filename_pattern='*wellplate.{well}_*height={z}mm*',z_ref=71.7,h_ref=110.2,filename_z_offset=-12))
    m=preview_names([name],s)['rows'][0]
    assert m['status']=='matched' and m['well']=='B1' and m['zg_mm']==124.8
    assert 'filename_z_offset' not in s
    expected=110.2+((124.8-12)-59.7)
    assert s['h_ref']+(m['zg_mm']-s['z_ref'])==pytest.approx(expected)
