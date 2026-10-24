import io
import math
import numpy as np
import pytest
from scipy.optimize import nnls
from openpyxl import load_workbook
from polar_core.science import owrk, _owrk_draw_fits
from tests.test_science import rows_for_energy
from tests.test_app import client


def test_batch_nnls_matches_scipy_including_boundaries():
    rng=np.random.default_rng(531)
    X=np.c_[np.ones(5),[0,.2,.7,1.,1.5]]
    ys=rng.uniform(0,15,(400,5))
    actual=_owrk_draw_fits(X,ys)
    expected=np.array([nnls(X,y)[0] for y in ys])
    np.testing.assert_allclose(actual,expected,atol=1e-12)
    assert np.any(actual[:,0]==0) and np.any(actual[:,1]==0)


def test_small_angle_error_matches_linear_propagation_and_total_covariance():
    rows=rows_for_energy()
    for row in rows:row['angle_error']=.1
    result=owrk(rows,'Test','static',20)
    u=result['uncertainty'];keys=['dispersive','polar','total']
    jacobian=[]
    for i in range(3):
        lo=[dict(r) for r in rows];hi=[dict(r) for r in rows]
        for row in lo+hi:row.pop('angle_error')
        lo[i]['angle']-=1e-4;hi[i]['angle']+=1e-4
        a=owrk(lo,'Test','static',20);b=owrk(hi,'Test','static',20)
        jacobian.append([(b[k]-a[k])/2e-4 for k in keys])
    predicted=np.linalg.norm(jacobian,axis=0)*.1
    np.testing.assert_allclose([u['errors'][k] for k in keys],predicted,rtol=.035)
    d,p,t=[u['errors'][k] for k in keys]
    assert t*t==pytest.approx(d*d+p*p+2*u['component_covariance'],rel=1e-10)
    assert abs(t-math.hypot(d,p))>.1*t
    assert result==owrk(rows,'Test','static',20)
    assert len(u['band'])==61


def test_missing_zero_and_partial_errors_are_distinct():
    rows=rows_for_energy()
    assert owrk(rows,'Test','static',20)['uncertainty'] is None
    rows[0]['angle_error']=0
    with pytest.raises(ValueError,match='every included measurement'):
        owrk(rows,'Test','static',20)
    for row in rows:row['angle_error']=0
    result=owrk(rows,'Test','static',20)
    assert result['uncertainty']['errors']==dict(total=0,dispersive=0,polar=0)
    assert all(p['y_error']==0 for p in result['data'])


@pytest.mark.parametrize('error',[-1,float('nan'),float('inf'),1000])
def test_invalid_angle_error_rejected(error):
    rows=rows_for_energy()
    for row in rows:row['angle_error']=error
    with pytest.raises(ValueError,match='[Aa]ngle error'):
        owrk(rows,'Test','static',20)


def test_replicates_do_not_artificially_reduce_directional_spread():
    rows=rows_for_energy()
    for row in rows:row.update(angle_error=2,error_kind='polar_spread')
    a=owrk(rows,'Test','static',20)
    b=owrk(rows*4,'Test','static',20)
    assert a['uncertainty']['errors']==b['uncertainty']['errors']
    assert all(p['angle_error']==2 and p['n']==4 for p in b['data'])
    assert any('uncertainties in mean' in w for w in b['warnings'])
    assert any('Between-drop scatter' in w for w in b['warnings'])


def test_bounded_large_error_and_boundary_fit_are_finite():
    rows=rows_for_energy()
    for row,angle,error in zip(rows,[179,30,120],[10,10,15]):row.update(angle=angle,angle_error=error)
    result=owrk(rows,'Test','static',20)
    u=result['uncertainty']
    assert u['boundary_fraction']>.5
    assert all(np.isfinite(v) and v>=0 for v in u['errors'].values())
    assert all(0<=p['low']<=p['high'] for p in u['band'])
    assert any('asymmetric' in w for w in result['warnings'])


@pytest.mark.parametrize('with_errors',[False,True])
def test_energy_export_matches_api_with_clean_columns(client,with_errors):
    rows=rows_for_energy()
    if with_errors:
        for row in rows:row['angle_error']=2
    payload=dict(rows=rows,substrate='Synthetic test',convention='static',temperature=20,confirmed=True)
    headers={'X-POLAR-Request':'1'}
    response=client.post('/api/owrk',json=payload,headers=headers)
    assert response.status_code==200,response.json
    result=response.json
    export=client.post('/api/owrk/export',json=payload,headers=headers)
    assert export.status_code==200
    wb=load_workbook(io.BytesIO(export.data));ws=wb.active
    assert wb.sheetnames==['Surface energy'] and ws.max_column==3 and ws.max_row==4
    assert '1 SD' in ws['C1'].value and ws['C1'].comment
    for i,key in enumerate(['total','dispersive','polar'],start=2):
        assert ws.cell(i,2).value==pytest.approx(result[key])
        if with_errors:assert ws.cell(i,3).value==pytest.approx(result['uncertainty']['errors'][key])
        else:assert ws.cell(i,3).value is None
