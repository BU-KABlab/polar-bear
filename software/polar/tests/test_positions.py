import io
import json
import time
import zipfile

import cv2
import numpy as np
import pytest
from openpyxl import load_workbook

from app import clean_settings,create_app
from polar_core.naming import filename_metadata,preview_names,select_positions
from tests.test_contact_angle import synthetic

PATTERN='camera_*_{position}_ursa_wellplate.{well}_after_dispense_white_height={z}mm*'
POSITIONS=['left_4mm','left_3mm','left_2mm','left_1mm','center_0mm',
           'right_1mm','right_2mm','right_3mm','right_4mm']

def settings():
    return clean_settings(dict(grouping='positions',filename_format='custom',filename_pattern=PATTERN))

def test_supplied_offset_names():
    names=['camera_176_left_4mm_ursa_wellplate.A1_after_dispense_white_height=124-200mm_b25_20261007-133305.tiff',
           'camera_204_left_1mm_ursa_wellplate.A1_after_dispense_white_height=124-600mm_b25_20261007-133435',
           'camera_219_center_0mm_ursa_wellplate.A1_after_dispense_white_height=123-600mm_b25_20261007-133519']
    result=preview_names(names,settings())
    assert result['matched']==3 and result['group_count']==3
    assert [(r['position'],r['well'],r['zg_mm']) for r in result['rows']]==[
        ('left_4mm','A1',124.2),('left_1mm','A1',124.6),('center_0mm','A1',123.6)]

@pytest.mark.parametrize('token,expected',[('=124-200',124.2),('_122p8',122.8),('_122P8',122.8),('_-1p2',-1.2),('_122.8',122.8)])
def test_mixed_height_styles(token,expected):
    s=clean_settings(dict(settings(),filename_pattern=PATTERN.replace('height=','height*')))
    name=f'camera_262_left_4mm_ursa_wellplate.A2_after_dispense_white_height{token}mm_20261007-135039.tiff'
    m=filename_metadata(name,s)
    assert m['zg_mm']==expected and m['position']=='left_4mm' and m['well']=='A2'

def test_wildcard_does_not_truncate_invalid_height():
    s=clean_settings(dict(settings(),filename_pattern=PATTERN.replace('height=','height*')))
    with pytest.raises(ValueError):filename_metadata('camera_262_left_4mm_ursa_wellplate.A2_after_dispense_white_height_122p8p2mm.tiff',s)

def test_position_selection_keeps_wells_and_failed_positions_separate():
    rows=[]
    for well,pos,target in [('A1','left_1mm',123.4),('A1','center_0mm',123.6),('B1','left_1mm',123.2)]:
        for z in [123.2,123.4,123.6]:
            rows.append(dict(well=well,position=pos,zg_mm=z,focus_score=3 if z==target else 1,
                             status='valid',filename=f'{well}-{pos}-{z}',brightness_group='b25',ellipse_axis_ratio=1))
    rows.append(dict(well='A1',position='right_4mm',zg_mm=123.2,status='rejected'))
    chosen={(r['well'],r['position']):r for r in select_positions(rows,.2)}
    assert len(chosen)==4
    assert chosen['A1','left_1mm']['zg_mm']==123.4
    assert chosen['A1','center_0mm']['zg_mm']==123.6
    assert chosen['B1','left_1mm']['zg_mm']==123.2
    assert chosen['A1','right_4mm']['status']=='failed'

def test_position_pattern_cannot_silently_merge_offsets():
    s=settings()
    for grouping in ['wells','timepoints']:
        with pytest.raises(ValueError):clean_settings(dict(s,grouping=grouping))
    with pytest.raises(ValueError,match='position'):
        clean_settings(dict(s,filename_pattern='*wellplate.{well}_height={z}mm*'))

def test_nine_positions_ten_heights_end_to_end(tmp_path,monkeypatch):
    monkeypatch.delenv('POLAR_API_TOKEN',raising=False)
    app=create_app(tmp_path,testing=True);client=app.test_client()
    try:
        image,_=synthetic();encoded={}
        for key,blur in [('sharp',.1),('blurred',1.5)]:
            _,png=cv2.imencode('.png',np.uint16(cv2.GaussianBlur(image,(0,0),blur)*65535))
            encoded[key]=png.tobytes()
        _,png=cv2.imencode('.png',np.zeros(image.shape,dtype=np.uint16));encoded['blank']=png.tobytes()
        archive=io.BytesIO();expected={}
        with zipfile.ZipFile(archive,'w') as z:
            for i,pos in enumerate(POSITIONS):
                best=1+i%8;expected[pos]=123.2+.2*best
                for h in range(10):
                    height=f'{123.2+.2*h:.3f}'.replace('.','-')
                    filename=f'camera_{i*10+h}_{pos}_ursa_wellplate.A1_after_dispense_white_height={height}mm_b25_20261007-133305.png'
                    key='blank' if i==8 else 'sharp' if h==best else 'blurred'
                    z.writestr('camera/images/'+filename,encoded[key])
        archive.seek(0)
        response=client.post('/api/jobs',data={'files':(archive,'offsets.zip'),'settings':json.dumps(settings())},headers={'X-POLAR-Request':'1'})
        assert response.status_code==202,response.json
        jid=response.json['id'];deadline=time.monotonic()+20
        while time.monotonic()<deadline:
            job=client.get('/api/jobs/'+jid).json
            if job['state'] not in ('queued','running'):break
            time.sleep(.03)
        assert job['state']=='complete',job['message']
        assert len(job['results'])==9 and len(job['candidates'])==90
        for r in job['results']:
            assert r['label']==f"A1 · {r['position']}"
            if r['position']=='right_4mm':assert r['angle'] is None
            else:
                assert r['zg']==pytest.approx(expected[r['position']])
                assert r['angle']>0
        wb=load_workbook(io.BytesIO(client.get(f'/api/jobs/{jid}/results.xlsx').data))
        assert wb.sheetnames==['Results'] and wb.active.max_column==3 and wb.active.max_row==10
        assert set(c.value for c in wb.active['A'][1:])=={f'A1 · {p}' for p in POSITIONS}
    finally:app.extensions['polar']['pool'].shutdown(wait=True)
