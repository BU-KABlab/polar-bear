import io
import json
import zipfile
import pytest
from werkzeug.datastructures import FileStorage
from openpyxl import load_workbook
from app import create_app,extract_uploads,excel_results

@pytest.fixture
def client(tmp_path,monkeypatch):
    monkeypatch.delenv('POLAR_API_TOKEN',raising=False)
    app=create_app(tmp_path,testing=True)
    yield app.test_client()
    app.extensions['polar']['pool'].shutdown(wait=True)

def test_local_api_guard(client):
    assert client.get('/').status_code==200
    assert client.get('/api/health').json['name']=='POLAR'
    assert client.post('/api/owrk',json={}).status_code==403
    assert client.post('/api/owrk',json={},headers={'X-POLAR-Request':'1'}).status_code==400
    assert client.get('/api/jobs',headers={'Origin':'https://example.com'}).status_code==403

def test_token(tmp_path,monkeypatch):
    monkeypatch.setenv('POLAR_API_TOKEN','test-secret')
    app=create_app(tmp_path,testing=True);c=app.test_client()
    assert c.get('/api/jobs').status_code==401
    assert c.get('/api/jobs',headers={'Authorization':'Bearer test-secret'}).status_code==200
    app.extensions['polar']['pool'].shutdown()

@pytest.mark.parametrize('name',['../evil.tiff','/tmp/evil.tiff','C:/evil.tiff'])
def test_zip_rejects_traversal(tmp_path,name):
    b=io.BytesIO()
    with zipfile.ZipFile(b,'w') as z:z.writestr(name,b'fake')
    b.seek(0)
    with pytest.raises(ValueError,match='unsafe'):extract_uploads([FileStorage(b,filename='input.zip')],tmp_path/'files')

def test_zip_nested_names_and_clean_excel(tmp_path):
    b=io.BytesIO()
    with zipfile.ZipFile(b,'w') as z:
        z.writestr('manifest.json','{}');z.writestr('images/camera_1_plate.C4_z=59-300mm.tiff',b'fake')
    b.seek(0)
    entries=extract_uploads([FileStorage(b,filename='input.zip')],tmp_path/'files')
    assert entries[0]['filename']=='camera_1_plate.C4_z=59-300mm.tiff'
    assert len(entries)==1
    wb=load_workbook(excel_results([dict(label='C4',angle=44.5,error=1.57),dict(label='=FORMULA()',angle=None,error=None)]))
    assert wb.sheetnames==['Results']
    ws=wb.active;assert ws.max_column==3
    assert [c.value for c in ws[1]]==['Well / frame','Contact angle (°)','Error (°)']
    assert ws['A3'].data_type=='s' and ws['B3'].value is None

def test_mixed_archive_rejected(tmp_path):
    with pytest.raises(ValueError,match='one ZIP'):
        extract_uploads([FileStorage(io.BytesIO(),filename='a.zip'),FileStorage(io.BytesIO(),filename='b.png')],tmp_path)

def test_large_multipart_zip_end_to_end(client):
    import cv2
    import numpy as np
    import time
    from tests.test_contact_angle import synthetic
    pixels,_=synthetic()
    ok,png=cv2.imencode('.png',np.uint8(np.clip(pixels,0,1)*255))
    assert ok
    b=io.BytesIO()
    with zipfile.ZipFile(b,'w') as z:
        z.writestr('manifest.json',' '*1024**2)
        z.writestr('images/camera_1_plate.C1_z=59-700mm_b25.png',png.tobytes())
    b.seek(0)
    response=client.post('/api/jobs',data={'files':(b,'scan.zip'),'settings':json.dumps({'volume':10})},headers={'X-POLAR-Request':'1'})
    assert response.status_code==202,response.json
    jid=response.json['id']
    for _ in range(100):
        job=client.get('/api/jobs/'+jid).json
        if job['state'] not in ('queued','running'):break
        time.sleep(.02)
    assert job['state']=='complete',job
    assert len(job['results'])==1
    assert job['results'][0]['angle'] is not None,job['results']
    assert client.get('/api/jobs/'+jid+'/results.xlsx').status_code==200

def test_outline_upload_and_confirmation(client):
    import cv2
    import numpy as np
    import time
    from tests.test_science import geometry
    r,b,H,R=geometry(60);scale=.005;yy,xx=np.mgrid[:1200,:1200]
    image=np.where((xx-600)**2+(yy-600)**2<(b/scale)**2,.2,.1)
    for a in np.arange(24)*2*np.pi/24:
        x=600+r/scale*np.cos(a);y=600+r/scale*np.sin(a)
        image+=.75*np.exp(-((xx-x)**2+(yy-y)**2)/(2*3**2))
    _,png=cv2.imencode('.png',(np.clip(image,0,1)*65535).astype(np.uint16))
    s=dict(mode='outline',branch='below',mm_per_pixel=scale,h_ref=H,z_ref=59.7)
    response=client.post('/api/jobs',data={'files':(io.BytesIO(png.tobytes()),'plate.C1_z=59-700mm.png'),'settings':json.dumps(s)},headers={'X-POLAR-Request':'1'})
    assert response.status_code==202
    jid=response.json['id']
    for _ in range(250):
        job=client.get('/api/jobs/'+jid).json
        if job['state'] not in ('queued','running'):break
        time.sleep(.02)
    assert job['state']=='complete'
    assert job['results'][0]['status']=='outline_review',job
    endpoint=f'/api/jobs/{jid}/outline/0';headers={'X-POLAR-Request':'1'}
    raw=dict(circle=dict(cx=600,cy=600,radius_px=b/scale))
    assert client.post(endpoint,json=raw,headers=headers).status_code==400
    raw['confirmed']=True
    answer=client.post(endpoint,json=raw,headers=headers)
    assert answer.status_code==200,answer.json
    assert answer.json['angle']==pytest.approx(60,abs=.2)
    assert answer.json['estimated_volume']>0
    assert any('boundary' in w for w in answer.json['warnings'])
    raw['circle']['radius_px']=800
    assert client.post(endpoint,json=raw,headers=headers).status_code==400

def test_unknown_volume_does_not_require_volume():
    from app import clean_settings
    assert clean_settings({'mode':'outline','volume':''})['volume'] is None

@pytest.mark.parametrize('grouping',['wells','frames'])
def test_elliptical_ring_returns_review_result_and_clean_export(client,grouping):
    import cv2
    import numpy as np
    import time
    from tests.test_contact_angle import synthetic
    pixels,_=synthetic()
    pixels=cv2.resize(pixels,None,fx=2.,fy=1.)
    _,png=cv2.imencode('.png',np.uint16(pixels*65535))
    response=client.post('/api/jobs',data={
        'files':(io.BytesIO(png.tobytes()),'plate.C1_z=59-700mm.png'),
        'settings':json.dumps({'volume':10,'grouping':grouping})},headers={'X-POLAR-Request':'1'})
    assert response.status_code==202
    jid=response.json['id']
    for _ in range(250):
        job=client.get('/api/jobs/'+jid).json
        if job['state'] not in ('queued','running'):break
        time.sleep(.02)
    assert job['state']=='complete'
    result=job['results'][0]
    assert result['angle']>0 and result['error']>0
    assert result['status']=='review' and result['provisional']
    assert len(result['pairs'])==12 and len(result['points'])==24
    assert any('Provisional result' in w for w in result['warnings'])
    audit=json.loads(client.get('/api/jobs/'+jid+'/audit.json').data)
    assert audit['results'][0]['quality']['ellipse_axis_ratio']>1.5
    wb=load_workbook(io.BytesIO(client.get('/api/jobs/'+jid+'/results.xlsx').data))
    assert wb.sheetnames==['Results'] and wb.active.max_column==3
    assert wb.active['B2'].value==pytest.approx(result['angle'])

def test_untrusted_host_rejected_without_token(client):
    assert client.get('/api/jobs',headers={'Host':'untrusted.example'}).status_code==400
