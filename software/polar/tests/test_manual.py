import io,json,time
import cv2
import numpy as np
import pytest
from openpyxl import load_workbook
from app import create_app,clean_settings
from polar_core.manual import measure_manual
from polar_core.contact_angle import solve_pair
from tests.test_contact_angle import synthetic

def pairs():
    _,pts=synthetic()
    return [[pts[i].tolist(),pts[i+12].tolist()] for i in range(12) if i!=7]

def test_manual_distances_match_solver_without_resizing():
    s=clean_settings({});lines=pairs();r=measure_manual(lines,700,360,s,123.6)
    expected=[solve_pair(np.linalg.norm(np.subtract(a,b))*s['mm_per_pixel']/2,29.75,111.8,10)['theta_deg'] for a,b in lines]
    assert r['angle']==pytest.approx(np.mean(expected))
    assert r['directional_sd']==pytest.approx(np.std(expected,ddof=1))
    assert r['error']**2==pytest.approx(r['directional_sd']**2+r['height_error']**2)
    assert len(r['pairs'])==11 and r['analysis_method']=='manual'
    assert r['manual_pairs']==lines and r['quality']['ellipse_axis_ratio']>1

def test_single_pair_does_not_invent_directional_uncertainty():
    r=measure_manual(pairs()[:1],700,360,clean_settings({}),123.6)
    assert r['angle']>0 and r['height_error']>0
    assert r['directional_sd'] is None and r['error'] is None
    json.dumps(r,allow_nan=False)

@pytest.mark.parametrize('lines',[[],[[[1,1],[1,1]]],[[[-1,2],[3,4]]],[[[0,0],[700,100]]],[[[1,2],[float('nan'),3]]],pairs()+pairs(),[pairs()[0],pairs()[0]]])
def test_manual_invalid_endpoints_rejected(lines):
    with pytest.raises(ValueError):measure_manual(lines,700,360,clean_settings({}),123.6)

def test_manual_outline_then_existing_boundary_solver():
    r=measure_manual(pairs(),700,360,clean_settings({'mode':'outline'}),123.6)
    assert r['status']=='outline_review' and r['angle'] is None
    assert len(r['points'])==22 and r['quality']['center_x_px']==pytest.approx(170)

def test_failed_measurement_can_be_manually_saved_and_exported(tmp_path,monkeypatch):
    monkeypatch.delenv('POLAR_API_TOKEN',raising=False)
    app=create_app(tmp_path,testing=True);c=app.test_client();headers={'X-POLAR-Request':'1'}
    try:
        image,_=synthetic(missing=7)
        _,png=cv2.imencode('.png',np.uint16(image*65535))
        response=c.post('/api/jobs',data={'files':[(io.BytesIO(png.tobytes()),f'plate.{well}_z=123-600mm.png') for well in ['A1','A2']],
                        'settings':json.dumps({'volume':10})},headers=headers)
        assert response.status_code==202
        jid=response.json['id']
        for _ in range(200):
            job=c.get(f'/api/jobs/{jid}').json
            if job['state']=='complete':break
            time.sleep(.02)
        assert job['results'][0]['angle'] is None
        sources=c.get(f'/api/jobs/{jid}/manual-sources/0').json['sources']
        assert len(sources)==1 and 'A1' in sources[0]['filename']
        source=sources[0]['file_id'];endpoint=f'/api/jobs/{jid}/manual-pairs/0'
        original=c.get(f'/api/jobs/{jid}/manual-image/0/{source}')
        assert original.status_code==200
        decoded=cv2.imdecode(np.frombuffer(original.data,np.uint8),cv2.IMREAD_UNCHANGED)
        assert decoded.shape[:2]==(360,700)
        assert c.post(endpoint,json={'file_id':'1','pairs':pairs()},headers=headers).status_code==400
        assert c.post(endpoint,json={'file_id':source,'pairs':[[[1,1],[1,1]]]},headers=headers).status_code==400
        assert c.get(f'/api/jobs/{jid}').json['results'][0]['angle'] is None
        saved=c.post(endpoint,json={'file_id':source,'pairs':pairs()},headers=headers)
        assert saved.status_code==200,saved.json
        r=saved.json;assert r['angle']>0 and r['status']=='review'
        assert r['manual_history'][0]['angle'] is None and len(r['manual_pairs'])==11
        assert c.get(f'/api/jobs/{jid}/preview/0').status_code==200
        audit=json.loads(c.get(f'/api/jobs/{jid}/audit.json').data)
        assert audit['results'][0]['analysis_method']=='manual'
        assert audit['results'][1]['angle'] is None
        wb=load_workbook(io.BytesIO(c.get(f'/api/jobs/{jid}/results.xlsx').data))
        assert wb.sheetnames==['Results'] and wb.active.max_column==3
        assert wb.active['B2'].value==pytest.approx(r['angle'])
        again=c.post(endpoint,json={'file_id':source,'pairs':pairs()[:1]},headers=headers)
        assert again.status_code==200 and again.json['error'] is None
        assert len(again.json['manual_history'])==2
    finally:app.extensions['polar']['pool'].shutdown(wait=True)
