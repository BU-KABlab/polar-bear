"""POLAR local analysis service. Start with: python app.py"""
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime,timezone
import io
import json
import math
import os
from pathlib import Path,PurePosixPath
import re
import secrets
import shutil
import stat
import threading
import tempfile
import zipfile

os.environ.setdefault('OPENCV_IO_MAX_IMAGE_PIXELS','40000000')
import cv2
import numpy as np
from flask import Flask,abort,jsonify,request,send_file,send_from_directory
from werkzeug.exceptions import HTTPException

from polar_core.contact_angle import read_image,detect_leds,analyze,ring_warnings
from polar_core.batch_contact_angle import focus_metrics,image_metadata,choose_images,height_error
from polar_core.naming import naming_settings,filename_metadata,preview_names,select_timepoints,select_positions
from polar_core.science import LIQUIDS,finite,fit_circle,propose_circles,outline_result,owrk
from polar_core.manual import measure_manual

ROOT=Path(__file__).resolve().parent
EXTENSIONS={'.tif','.tiff','.png','.jpg','.jpeg','.bmp'}


def clean_settings(raw):
    mode=raw.get('mode','volume');grouping=raw.get('grouping','wells')
    if mode not in ('volume','outline') or grouping not in ('wells','frames','timepoints','positions'): raise ValueError('Invalid analysis mode.')
    settings={'mode':mode,'grouping':grouping,'volume':finite(raw.get('volume',10),'Volume',.001,100000) if mode=='volume' else None,
              'height_step':finite(raw.get('height_step',.2),'Height step',.0001,10),
              'mm_per_pixel':finite(raw.get('mm_per_pixel',6.35/909.85),'Pixel scale',1e-7,10),
              'led_spacing':finite(raw.get('led_spacing',59.5),'LED spacing',.01,1000),
              'z_ref':finite(raw.get('z_ref',123.5),'Reference Z',-10000,10000),
              'h_ref':finite(raw.get('h_ref',111.7),'Reference LED height',.001,10000),
              'substrate':str(raw.get('substrate',''))[:160].strip(),
              'liquid':str(raw.get('liquid',''))[:80].strip(),
              'branch':raw.get('branch','below'),'boundary':raw.get('boundary','silhouette'),
              'convention':raw.get('convention','static')}
    if settings['branch'] not in ('below','above') or settings['boundary'] not in ('silhouette','contact'):
        raise ValueError('Invalid outline interpretation.')
    if settings['convention'] not in ('static','advancing','receding','apparent'): raise ValueError('Invalid angle convention.')
    settings['zg_fallback']=finite(raw['zg_fallback'],'Fallback Z',-10000,10000) if raw.get('zg_fallback') not in ('',None) else None
    settings.update(naming_settings({**raw,**settings}))
    return settings


def extract_uploads(files,destination):
    destination.mkdir(parents=True,exist_ok=True)
    entries=[];total=0;seen=set()
    if any((f.filename or '').lower().endswith('.zip') for f in files) and len(files)!=1:
        raise ValueError('Upload one ZIP per run, or select individual images. Separate archives may reuse well names.')
    def store(name,source,size):
        nonlocal total
        if len(entries)>=1000 or size>128*1024**2: raise ValueError('Limit: 1,000 images and 128 MB per image.')
        if Path(name).suffix.lower() not in EXTENSIONS: return
        total+=size
        if total>4*1024**3: raise ValueError('Extracted images exceed 4 GB.')
        base=Path(name.replace('\\','/')).name
        base=re.sub(r'[\x00-\x1f]','',base)[:200]
        saved=f'{len(entries):04d}__{base}'
        with (destination/saved).open('wb') as target:
            copied=0
            while True:
                block=source.read(1024*1024)
                if not block: break
                copied+=len(block)
                if copied>128*1024**2 or copied>size: raise ValueError('Image exceeds its declared size.')
                target.write(block)
        entries.append({'stored':saved,'filename':base,'archive_path':name})
    for upload in files:
        filename=upload.filename or ''
        if filename.lower().endswith('.zip'):
            try:
                # Python 3.10's rolled upload stream lacks ZipFile's seekable API.
                with tempfile.TemporaryFile() as zip_stream:
                    shutil.copyfileobj(upload.stream,zip_stream)
                    zip_stream.seek(0)
                    with zipfile.ZipFile(zip_stream) as archive:
                        members=archive.infolist()
                        if len(members)>5000: raise ValueError('Archive contains too many entries.')
                        for member in members:
                            name=member.filename.replace('\\','/')
                            path=PurePosixPath(name)
                            if not path.parts or path.is_absolute() or '..' in path.parts or ':' in path.parts[0] or stat.S_ISLNK(member.external_attr>>16):
                                raise ValueError('Archive contains an unsafe path or symbolic link.')
                            if member.is_dir() or '__MACOSX' in path.parts or path.name.startswith('._'): continue
                            if name in seen: raise ValueError('Archive contains duplicate file paths.')
                            seen.add(name)
                            if member.flag_bits&1: raise ValueError('Encrypted archives are not supported.')
                            if Path(name).suffix.lower() in EXTENSIONS:
                                with archive.open(member) as src: store(name,src,member.file_size)
            except zipfile.BadZipFile as error: raise ValueError('The ZIP archive is damaged or unreadable.') from error
        else:
            upload.stream.seek(0,2);size=upload.stream.tell();upload.stream.seek(0)
            store(filename,upload.stream,size)
    if not entries: raise ValueError('No supported images found. Upload a CubOS ZIP or TIFF/PNG/JPEG images.')
    return entries


def excel_results(results,title='POLAR results',timepoints=False):
    # This is the portable application's runtime exporter, independent of Codex tooling.
    from openpyxl import Workbook
    from openpyxl.styles import Font,PatternFill,Alignment
    wb=Workbook();ws=wb.active;ws.title='Results';ws.sheet_view.showGridLines=False
    ws.append(['Well','Time (s)','Contact angle (°)','Error (°)'] if timepoints else ['Well / frame','Contact angle (°)','Error (°)'])
    for result in results:
        ws.append([result.get('well',result['label']),result.get('elapsed_s'),result.get('angle'),result.get('error')] if timepoints else [result['label'],result.get('angle'),result.get('error')])
        ws.cell(ws.max_row,1).data_type='s'
    for cell in ws[1]:
        cell.font=Font(name='Arial',bold=True,color='FFFFFF',size=11)
        cell.fill=PatternFill('solid',fgColor='173F46');cell.alignment=Alignment(vertical='center')
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font=Font(name='Arial',size=11)
            if cell.column>1:cell.number_format='0.000'
            if cell.row%2==0:cell.fill=PatternFill('solid',fgColor='F0F6F6')
    ws.column_dimensions['A'].width=30;ws.column_dimensions['B'].width=25;ws.column_dimensions['C'].width=20
    if timepoints:ws.column_dimensions['D'].width=20
    ws.freeze_panes='A2';ws.row_dimensions[1].height=28
    ws.auto_filter.ref=f"A1:{'D' if timepoints else 'C'}{ws.max_row}"
    wb.properties.title=title
    wb.properties.description='Error: directional SD + height-step contribution in quadrature (k=1); not total uncertainty. Blank results require review in POLAR.'
    output=io.BytesIO();wb.save(output);output.seek(0);return output


def create_app(data_dir=None,testing=False):
    app=Flask(__name__,static_folder=str(ROOT/'static'))
    app.config.update(MAX_CONTENT_LENGTH=512*1024**2,TESTING=testing)
    data=Path(data_dir or os.environ.get('POLAR_DATA_DIR',ROOT/'data')).resolve();data.mkdir(parents=True,exist_ok=True)
    lock=threading.RLock();pool=ThreadPoolExecutor(max_workers=1,thread_name_prefix='polar')
    token=os.environ.get('POLAR_API_TOKEN','')
    if not token:
        app.config['TRUSTED_HOSTS']=['localhost','127.0.0.1','[::1]']
    def folder(jid):
        if not re.fullmatch(r'[a-f0-9]{16}',jid): abort(404)
        p=data/jid
        if not p.is_dir():abort(404)
        return p
    def load(jid):
        with lock:
            p=folder(jid)/'job.json'
            if not p.exists():abort(404)
            return json.loads(p.read_text())
    def save(job):
        with lock:
            p=data/job['id'];p.mkdir(exist_ok=True)
            temp=p/'job.tmp';temp.write_text(json.dumps(job,allow_nan=False));temp.replace(p/'job.json')
    def manual_context(jid,rid):
        job=load(jid)
        if job['state']!='complete':raise ValueError('Wait for analysis to finish before marking LED pairs.')
        result=next((r for r in job['results'] if r['id']==rid),None)
        if result is None:abort(404)
        sources=[];settings=job['settings']
        for index,entry in enumerate(job['files']):
            try:meta=filename_metadata(entry['filename'],settings)
            except ValueError:continue
            if settings['grouping']=='frames':
                if entry['filename']!=result['filename']:continue
            else:
                if meta['well']!=result.get('well'):continue
                if settings['grouping']=='positions' and meta.get('position')!=result.get('position'):continue
                if settings['grouping']=='timepoints' and meta.get('elapsed_s')!=result.get('elapsed_s'):continue
            sources.append(dict(file_id=str(index),filename=entry['filename'],stored=entry['stored'],zg=meta['zg_mm']))
        return job,result,sources
    def manual_source(jid,rid,fid):
        job,result,sources=manual_context(jid,rid)
        source=next((s for s in sources if s['file_id']==str(fid)),None)
        if source is None:raise ValueError('Select an image belonging to this measurement group.')
        return job,result,source
    if not testing:
        for p in data.glob('*/job.json'):
            job=json.loads(p.read_text())
            if job['state'] in ('queued','running'):
                job.update(state='interrupted',message='Analysis was interrupted. Re-upload to start a new run.');save(job)
    def process(jid):
        job=load(jid);settings=job['settings'];base=folder(jid)
        try:
            job.update(state='running',message='Finding LED rings',progress=0);save(job)
            candidates=[];mapping={}
            for i,entry in enumerate(job['files']):
                path=base/'images'/entry['stored'];mapping[str(path)]=entry
                row={'filename':entry['filename'],'path':str(path),'status':'rejected','reason':''}
                try:
                    row.update(filename_metadata(entry['filename'],settings))
                    gray,_=read_image(path);spots,q=detect_leds(gray)
                    row.update(focus_metrics(gray,spots));row.update(ellipse_axis_ratio=q['ellipse_axis_ratio'],status='valid')
                except (ValueError,OSError,cv2.error) as error:row['reason']=str(error)
                candidates.append(row)
                job.update(progress=round(75*(i+1)/len(job['files'])),message=f"Scanned {i+1} of {len(job['files'])} images");save(job)
            # CubOS timestamps preserve chronology even when numeric image IDs change width.
            if settings['grouping']=='frames' and all(re.search(r'_(\d{8}-\d{6})',c['filename']) for c in candidates):
                candidates.sort(key=lambda c:re.search(r'_(\d{8}-\d{6})',c['filename']).group(1))
            if settings['grouping']=='timepoints':
                selected=select_timepoints(candidates,settings['height_step'])
            elif settings['grouping']=='positions':
                selected=select_positions(candidates,settings['height_step'])
            elif settings['grouping']=='wells':
                selected=choose_images(candidates,settings['height_step'])
            else:
                selected=[]
                for row in candidates:
                    if row['status']=='valid':selected.append(dict(row,status='review',warnings=['Single frame; focus selection not evaluated.']))
                    else:selected.append(dict(row,status='failed',warnings=[row['reason']],well=row.get('well',row['filename'])))
            results=[]
            for i,item in enumerate(selected):
                rid=str(i)
                label=item['well'] if settings['grouping'] in ('wells','timepoints','positions') else item['filename']
                if settings['grouping']=='positions':label+=f" · {item['position']}"
                if settings['grouping']=='timepoints':
                    label+=f" · {item['elapsed_s']:g} s"
                result={'id':rid,'label':label,'well':item.get('well'),'filename':item.get('filename',''),
                        'angle':None,'error':None,'warnings':item.get('warnings',[]),'status':'failed',
                        'substrate':settings['substrate'],'liquid':settings['liquid'],'convention':settings['convention']}
                if 'elapsed_s' in item:result['elapsed_s']=item['elapsed_s']
                if 'position' in item:result['position']=item['position']
                if item['status']!='failed':
                    path=Path(item['path']);zg=item['zg_mm'];entry=mapping[str(path)]
                    result.update(zg=zg,stored=entry['stored'],focus_score=item['focus_score'])
                    try:
                        gray,display=read_image(path);points,q=detect_leds(gray)
                        result.update(points=points.tolist(),quality=q,provisional=q['provisional'],image_width=gray.shape[1],image_height=gray.shape[0])
                        for warning in ring_warnings(q['ellipse_axis_ratio']):
                            if warning not in result['warnings']:result['warnings'].append(warning)
                        display_width=min(1500,gray.shape[1]);scaled=cv2.resize(display,(display_width,round(gray.shape[0]*display_width/gray.shape[1])))
                        cv2.imwrite(str(base/f'preview_{rid}.jpg'),scaled)
                        if settings['mode']=='volume':
                            solved,_=analyze(path,volume=settings['volume'],zg=zg,z_ref=settings['z_ref'],h_ref=settings['h_ref'],
                                             spacing=settings['led_spacing'],mm_per_pixel=settings['mm_per_pixel'])
                            uncertainty=height_error(solved,settings['height_step'])
                            result.update(angle=solved['contact_angle_mean_deg'],directional_sd=solved['contact_angle_sample_std_deg'],
                                          error=uncertainty['combined_direction_height_deg'],height_error=uncertainty['height_angle_standard_uncertainty_deg'],
                                          pairs=solved['pairs'],status='review' if result['warnings'] else 'complete',supplied_volume=settings['volume'])
                        else:
                            center=np.array([q['center_x_px'],q['center_y_px']])
                            result.update(circles=propose_circles(path,center,min(q['ellipse_diameters_px'])/2),status='outline_review')
                    except (ValueError,OSError,cv2.error) as error:result['warnings'].append(str(error))
                results.append(result)
                job.update(results=results,progress=75+round(25*(i+1)/max(1,len(selected))),message=f'Analyzing selection {i+1} of {len(selected)}');save(job)
            job.update(state='complete',progress=100,results=results,candidates=candidates,
                       message='Analysis complete' if results else 'No usable wells. Check filenames and the import details.')
            save(job)
        except Exception as error:
            app.logger.exception('Analysis failed')
            job.update(state='failed',message=str(error));save(job)
    @app.before_request
    def protect():
        if request.path.startswith('/api/') and request.path!='/api/health':
            if token and not secrets.compare_digest(request.headers.get('Authorization',''),'Bearer '+token):
                return jsonify(error='Enter the configured POLAR API token in Connection settings.'),401
            origin=request.headers.get('Origin')
            if origin and origin!=request.host_url.rstrip('/'):
                return jsonify(error='Cross-origin requests are not allowed.'),403
            if request.method in ('POST','PUT','DELETE') and not token and request.headers.get('X-POLAR-Request')!='1':
                return jsonify(error='Missing local request header.'),403
    @app.after_request
    def headers(response):
        response.headers['X-Content-Type-Options']='nosniff'
        response.headers['Content-Security-Policy']="default-src 'self'; img-src 'self' blob: data:; style-src 'self' 'unsafe-inline'; script-src 'self'; connect-src 'self'; frame-ancestors 'none'"
        response.headers['Cache-Control']='no-store' if request.path.startswith('/api/') else 'no-cache'
        return response
    @app.errorhandler(Exception)
    def handle(error):
        if isinstance(error,HTTPException):return jsonify(error=error.description),error.code
        if isinstance(error,(ValueError,KeyError,TypeError)):return jsonify(error=str(error)),400
        app.logger.exception('Request failed');return jsonify(error='Server error. See the local POLAR log.'),500
    @app.get('/')
    def home():return send_from_directory(app.static_folder,'index.html')
    @app.get('/api/health')
    def health():return jsonify(name='POLAR',version='0.6.0',auth_required=bool(token))
    @app.post('/api/naming/preview')
    def naming_preview():
        raw=request.json
        return jsonify(preview_names(raw.get('filenames',[]),clean_settings({**raw.get('settings',{}),'volume':10})))
    @app.get('/api/liquids')
    def liquids():return jsonify(LIQUIDS)
    @app.get('/api/jobs')
    def jobs():
        with lock:
            all_jobs=[json.loads(p.read_text()) for p in data.glob('*/job.json')]
        return jsonify([{k:j.get(k) for k in ('id','name','created','state','settings','results','message')} for j in sorted(all_jobs,key=lambda j:j['created'],reverse=True)])
    @app.post('/api/jobs')
    def upload():
        active=sum(json.loads(p.read_text())['state'] in ('queued','running') for p in data.glob('*/job.json'))
        if active>=3: return jsonify(error='Three runs are already queued. Wait for a run to finish.'),429
        settings=clean_settings(json.loads(request.form.get('settings','{}')))
        files=request.files.getlist('files')
        if not files:raise ValueError('Choose at least one ZIP or image.')
        jid=secrets.token_hex(8);base=data/jid
        try:
            entries=extract_uploads(files,base/'images')
            if settings['filename_format']=='custom':
                parsed=preview_names([e['filename'] for e in entries],settings)
                bad=[r for r in parsed['rows'] if r['status']=='unmatched']
                if bad:
                    raise ValueError(f"{len(bad)} filename(s) do not match your naming settings. First: {bad[0]['filename']} — {bad[0]['reason']}")
                if settings['grouping']=='wells':
                    by_well={}
                    for row in parsed['rows']:
                        if 'elapsed_s' in row:by_well.setdefault(row['well'],set()).add(row['elapsed_s'])
                    if any(len(times)>1 for times in by_well.values()):
                        raise ValueError('Multiple time points found within a well. Choose Time series · focus stack per time point, or Individual frames.')
        except Exception:shutil.rmtree(base,ignore_errors=True);raise
        job={'id':jid,'name':str(request.form.get('name') or Path(files[0].filename).stem)[:160],
             'created':datetime.now(timezone.utc).isoformat(),'state':'queued','progress':0,'message':'Waiting to analyze',
             'settings':settings,'files':entries,'results':[]}
        save(job);pool.submit(process,jid);return jsonify(job),202
    @app.get('/api/jobs/<jid>')
    def get_job(jid):return jsonify(load(jid))
    @app.get('/api/jobs/<jid>/preview/<rid>')
    def preview(jid,rid):
        if not rid.isdigit():abort(404)
        return send_file(folder(jid)/f'preview_{rid}.jpg')
    @app.post('/api/circle-fit')
    def circle_fit():return jsonify(fit_circle(request.json.get('points',[])))
    @app.get('/api/jobs/<jid>/manual-sources/<rid>')
    def manual_sources(jid,rid):
        _,_,sources=manual_context(jid,rid)
        return jsonify(sources=sources)
    @app.get('/api/jobs/<jid>/manual-image/<rid>/<fid>')
    def manual_image(jid,rid,fid):
        _,_,source=manual_source(jid,rid,fid)
        _,display=read_image(folder(jid)/'images'/source['stored'])
        ok,encoded=cv2.imencode('.png',display)
        if not ok:raise ValueError('Could not render the source image.')
        return send_file(io.BytesIO(encoded.tobytes()),mimetype='image/png')
    @app.post('/api/jobs/<jid>/manual-pairs/<rid>')
    def manual_pairs(jid,rid):
        with lock:
            payload=request.json or {}
            job,result,source=manual_source(jid,rid,payload.get('file_id'))
            gray,display=read_image(folder(jid)/'images'/source['stored'])
            measured=measure_manual(payload.get('pairs'),gray.shape[1],gray.shape[0],job['settings'],source['zg'])
            history=result.get('manual_history',[])+[{k:v for k,v in result.items() if k!='manual_history'}]
            identity={k:result[k] for k in ('id','label','well','position','elapsed_s','substrate','liquid','convention') if k in result}
            result.clear();result.update(identity,**measured,filename=source['filename'],stored=source['stored'],
                                          manual_history=history,manual_saved_utc=datetime.now(timezone.utc).isoformat())
            width=min(1500,gray.shape[1])
            preview=cv2.resize(display,(width,round(gray.shape[0]*width/gray.shape[1])))
            if not cv2.imwrite(str(folder(jid)/f'preview_{rid}.jpg'),preview):raise ValueError('Could not save the preview.')
            save(job);return jsonify(result)
    @app.post('/api/jobs/<jid>/outline/<rid>')
    def confirm_outline(jid,rid):
        with lock:
            job=load(jid)
            if job['state']!='complete' or job['settings']['mode']!='outline':raise ValueError('This run is not ready for outline analysis.')
            result=next((r for r in job['results'] if r['id']==rid),None)
            if not result or not result.get('points'):raise ValueError('No valid LED reflections for this image.')
            if request.json.get('confirmed') is not True:raise ValueError('Review and confirm the physical droplet boundary.')
            raw=request.json.get('circle',{})
            c={'cx':finite(raw.get('cx'),'Circle x',0,result['image_width']),
               'cy':finite(raw.get('cy'),'Circle y',0,result['image_height']),
               'radius_px':finite(raw.get('radius_px'),'Circle radius',1,max(result['image_width'],result['image_height']))}
            for key in ('relative_rms','angular_coverage_deg'):
                if key in raw:c[key]=finite(raw[key],key,0)
            if c.get('relative_rms',0)>.06:raise ValueError('Boundary points are not sufficiently circular for this model.')
            if c.get('angular_coverage_deg',360)<180:raise ValueError('Spread boundary points over at least half the droplet outline.')
            x,y,rad=c['cx'],c['cy'],c['radius_px']
            if x-rad<0 or y-rad<0 or x+rad>result['image_width'] or y+rad>result['image_height']:
                raise ValueError('The fitted circle is clipped by the image. Use a complete outline.')
            center=np.array([result['quality']['center_x_px'],result['quality']['center_y_px']])
            if np.linalg.norm(np.array([x,y])-center)>.2*rad:raise ValueError('Outline and LED centers disagree. Check the boundary or alignment.')
            settings=dict(job['settings']);settings['height']=settings['h_ref']+(result['zg']-settings['z_ref'])
            solved=outline_result(np.array(result['points']),c,settings,settings['height_step'])
            solved['warnings']=list(dict.fromkeys(result.get('warnings',[])+solved['warnings']))
            result.update(solved,status='review',outline_confirmed=True)
            save(job);return jsonify(result)
    @app.get('/api/jobs/<jid>/results.xlsx')
    def xlsx(jid):
        j=load(jid)
        if j['state']!='complete':raise ValueError('Wait for the analysis to complete.')
        return send_file(excel_results(j['results'],timepoints=j['settings']['grouping']=='timepoints'),download_name='POLAR_results.xlsx',as_attachment=True,
                         mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    @app.get('/api/jobs/<jid>/audit.json')
    def audit(jid):return send_file(folder(jid)/'job.json',as_attachment=True,download_name='POLAR_analysis.json')
    @app.post('/api/owrk')
    def surface_energy():
        payload=request.json
        if payload.get('confirmed') is not True:raise ValueError('Confirm substrate consistency and liquid properties before fitting.')
        return jsonify(owrk(payload.get('rows'),payload.get('substrate',''),payload.get('convention','static'),payload.get('temperature',20)))
    @app.post('/api/owrk/export')
    def export_surface():
        from openpyxl import Workbook
        p=request.json
        if p.get('confirmed') is not True:raise ValueError('Confirm the input data.')
        r=owrk(p.get('rows'),p.get('substrate',''),p.get('convention','static'),p.get('temperature',20))
        wb=Workbook();ws=wb.active;ws.title='Surface energy';ws.append(['Component','Surface energy (mJ/m²)','Propagated error · 1 SD (mJ/m²)'])
        for field in ('total','dispersive','polar'):ws.append([field.capitalize(),r[field],r['uncertainty']['errors'][field] if r['uncertainty'] else None])
        ws.column_dimensions['A'].width=20;ws.column_dimensions['B'].width=28;ws.column_dimensions['C'].width=36
        from openpyxl.comments import Comment
        from openpyxl.styles import Font,PatternFill
        ws['C1'].comment=Comment((r['uncertainty']['scope']+' '+r['uncertainty']['distribution']+' '+r['uncertainty']['replicate_policy']) if r['uncertainty'] else 'Angle errors were not supplied. Blank errors are unavailable, not zero.','POLAR')
        ws.freeze_panes='A2'
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='23716D')
        for row in ws.iter_rows(min_row=2,min_col=2):
            for cell in row:cell.number_format='0.00'
        wb.properties.description=f"Measurement temperature: {r['temperature_c']} °C (recorded only). {r['method']}"
        output=io.BytesIO();wb.save(output);output.seek(0)
        return send_file(output,as_attachment=True,download_name='POLAR_surface_energy.xlsx')
    app.extensions['polar']={'data':data,'pool':pool,'process':process,'save':save,'load':load}
    return app


if __name__=='__main__':
    host=os.environ.get('POLAR_HOST','127.0.0.1')
    if host not in ('127.0.0.1','localhost','::1') and not os.environ.get('POLAR_API_TOKEN'):
        raise SystemExit('Set POLAR_API_TOKEN before enabling network access.')
    app=create_app()
    app.run(host=host,port=int(os.environ.get('POLAR_PORT','8765')),debug=False,threaded=True)
